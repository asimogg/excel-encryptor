"""Excel Encryptor testleri — şifreleme turu, tip korunumu ve güvenlik kuralları.

Çalıştırma:  pytest -q
"""

from __future__ import annotations

import datetime as dt
import json
import os
import subprocess
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import crypto_engine as ce      # noqa: E402
import excel_io as xio          # noqa: E402

GOOD_KEY = "Gizli-Anahtar-2026"
WRONG_KEY = "Baska-Anahtar-2026"


# --------------------------------------------------------------------------- #
# Yardımcılar
# --------------------------------------------------------------------------- #

@pytest.fixture
def workbook(tmp_path):
    """Her tipten veri içeren örnek bir Excel dosyası."""
    path = tmp_path / "Personel Listesi.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personel"
    ws.append(["Ad Soyad", "TC Kimlik No", "Maaş", "İşe Giriş", "Aktif", "Not"])
    ws.append(["Ayşe Yılmaz", 12345678901, 84500.75, dt.datetime(2021, 3, 1), True, "ör-ünïcode ✓"])
    ws.append(["Mehmet Öz", 10987654321, 92000, dt.datetime(2019, 11, 15), False, None])
    ws.append(["Zeynep K.", 55544433322, 71250.5, dt.datetime(2023, 7, 4), True, "x"])
    w2 = wb.create_sheet("Formüller")
    w2.append(["Kod", "Oran"])
    w2.append(["A-1", 0.42])
    wb.save(path)
    return str(path)


def snapshot(path):
    wb = openpyxl.load_workbook(path)
    return {n: [[c.value for c in r] for r in wb[n].iter_rows()]
            for n in xio.data_sheet_names(wb)}


SELECTION = {
    "Personel": {(r, c) for r in range(2, 5) for c in range(1, 7)},
    "Formüller": {(2, 2)},
}


# --------------------------------------------------------------------------- #
# Şifrele → çöz turu
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("method", [m.id for m in ce.METHODS])
def test_roundtrip_preserves_every_value_and_type(workbook, method):
    original = snapshot(workbook)

    enc, n = xio.encrypt_workbook(
        workbook, {k: set(v) for k, v in SELECTION.items()}, method, GOOD_KEY)
    assert n == 18, "boş hücre (None) şifrelenmemeli (19 seçili, 1 tanesi boş)"
    assert os.path.basename(enc) == "Personel Listesi_encrypted.xlsx"

    dec, n2 = xio.decrypt_workbook(enc, GOOD_KEY)
    assert os.path.basename(dec) == "Personel Listesi_decrypted.xlsx"
    assert n2 == n
    assert snapshot(dec) == original, "sayı/tarih/bool/metin birebir geri gelmeli"


@pytest.mark.parametrize("method", [m.id for m in ce.METHODS])
def test_no_plaintext_survives_in_encrypted_file(workbook, method):
    enc, _ = xio.encrypt_workbook(
        workbook, {k: set(v) for k, v in SELECTION.items()}, method, GOOD_KEY)

    # Hücre değerlerinde de ham dosya baytlarında da açık veri kalmamalı
    text = str(snapshot(enc))
    raw = open(enc, "rb").read()
    for secret in ("Ayşe Yılmaz", "12345678901", "84500.75"):
        assert secret not in text
        assert secret.encode("utf-8") not in raw


@pytest.mark.parametrize("method", [m.id for m in ce.METHODS])
def test_wrong_key_is_rejected(workbook, method):
    enc, _ = xio.encrypt_workbook(
        workbook, {k: set(v) for k, v in SELECTION.items()}, method, GOOD_KEY)
    with pytest.raises(xio.ExcelError, match="Anahtar hatalı"):
        xio.decrypt_workbook(enc, WRONG_KEY)


def test_deterministic_method_maps_equal_values_to_equal_ciphertext(workbook):
    wb = openpyxl.load_workbook(workbook)
    wb["Personel"].cell(row=2, column=6).value = "TEKRAR"
    wb["Personel"].cell(row=4, column=6).value = "TEKRAR"
    wb.save(workbook)

    enc, _ = xio.encrypt_workbook(
        workbook, {"Personel": {(2, 6), (4, 6)}}, "aesgcm-det", GOOD_KEY)
    got = openpyxl.load_workbook(enc)["Personel"]
    assert got.cell(row=2, column=6).value == got.cell(row=4, column=6).value


def test_random_method_maps_equal_values_to_different_ciphertext(workbook):
    wb = openpyxl.load_workbook(workbook)
    wb["Personel"].cell(row=2, column=6).value = "TEKRAR"
    wb["Personel"].cell(row=4, column=6).value = "TEKRAR"
    wb.save(workbook)

    enc, _ = xio.encrypt_workbook(
        workbook, {"Personel": {(2, 6), (4, 6)}}, "aesgcm", GOOD_KEY)
    got = openpyxl.load_workbook(enc)["Personel"]
    assert got.cell(row=2, column=6).value != got.cell(row=4, column=6).value


def test_tampered_ciphertext_is_detected(workbook):
    """AEAD: şifreli metin değiştirilirse sessizce yanlış veri dönmemeli."""
    enc, _ = xio.encrypt_workbook(
        workbook, {"Personel": {(2, 1)}}, "aesgcm", GOOD_KEY)
    wb = openpyxl.load_workbook(enc)
    token = wb["Personel"].cell(row=2, column=1).value
    prefix, method, payload = token.split(":", 2)
    flipped = ("B" if payload[5] != "B" else "C")
    wb["Personel"].cell(row=2, column=1).value = (
        f"{prefix}:{method}:{payload[:5]}{flipped}{payload[6:]}")
    wb.save(enc)

    with pytest.raises(xio.ExcelError):
        xio.decrypt_workbook(enc, GOOD_KEY)


def test_output_never_overwrites_existing_file(workbook):
    first, _ = xio.encrypt_workbook(
        workbook, {"Personel": {(2, 1)}}, "aesgcm", GOOD_KEY)
    second, _ = xio.encrypt_workbook(
        workbook, {"Personel": {(2, 2)}}, "aesgcm", GOOD_KEY)
    assert first != second
    assert os.path.exists(first) and os.path.exists(second)


def test_source_file_is_left_untouched(workbook):
    before = open(workbook, "rb").read()
    xio.encrypt_workbook(workbook, {"Personel": {(2, 1)}}, "aesgcm", GOOD_KEY)
    assert open(workbook, "rb").read() == before


# --------------------------------------------------------------------------- #
# Güvenlik kuralları
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("weak", ["xor", "vigenere", "base64", "caesar", "rot13"])
def test_insecure_methods_are_not_offered(weak):
    """Gerçek koruma sağlamayan yöntemler kataloğa geri sızmamalı."""
    assert weak not in ce.METHODS_BY_ID


def test_every_method_is_authenticated_encryption(workbook):
    """Her yöntem kurcalamayı yakalayabilmeli (AEAD)."""
    for method in ce.METHODS:
        enc, _ = xio.encrypt_workbook(
            workbook, {"Personel": {(2, 1)}}, method.id, GOOD_KEY)
        wb = openpyxl.load_workbook(enc)
        token = wb["Personel"].cell(row=2, column=1).value
        p, m, payload = token.split(":", 2)
        wb["Personel"].cell(row=2, column=1).value = (
            f"{p}:{m}:{payload[:-6]}{'A' if payload[-6] != 'A' else 'B'}{payload[-5:]}")
        wb.save(enc)
        with pytest.raises(xio.ExcelError):
            xio.decrypt_workbook(enc, GOOD_KEY)
        os.remove(enc)


def test_kdf_meets_owasp_minimum():
    assert ce.KDF_ITERATIONS >= 600_000


@pytest.mark.parametrize("bad", ["", "kisa", "parolaparolaparola", "1234567"])
def test_weak_passwords_are_refused(bad):
    assert ce.password_problem(bad) is not None


def test_strong_password_is_accepted():
    assert ce.password_problem(GOOD_KEY) is None


def test_empty_password_cannot_derive_a_key():
    with pytest.raises(ce.CryptoError):
        ce.derive_key("", ce.new_salt())


def test_metadata_leaks_nothing_beyond_what_decryption_needs(workbook):
    enc, _ = xio.encrypt_workbook(
        workbook, {k: set(v) for k, v in SELECTION.items()}, "aesgcm", GOOD_KEY)
    meta = json.loads(openpyxl.load_workbook(enc)[ce.META_SHEET].cell(1, 1).value)

    assert set(meta) == {"version", "method", "kdf", "iterations", "salt"}
    # anahtar doğrulama etiketi = çevrimdışı parola deneme oracle'ı
    assert "verifier" not in meta
    # hangi hücrelerin gizlendiği bilgisi = hedef haritası
    assert "cells" not in meta
    assert GOOD_KEY not in json.dumps(meta)


def test_salt_is_unique_per_file(workbook):
    salts = set()
    for _ in range(3):
        enc, _ = xio.encrypt_workbook(
            workbook, {"Personel": {(2, 1)}}, "aesgcm", GOOD_KEY)
        meta = json.loads(openpyxl.load_workbook(enc)[ce.META_SHEET].cell(1, 1).value)
        salts.add(meta["salt"])
        os.remove(enc)
    assert len(salts) == 3


def test_xml_parser_is_hardened():
    """openpyxl güvenilmeyen XML'e karşı korumasız; defusedxml devrede olmalı."""
    assert xio.defusedxml is not None


def test_unknown_method_in_file_is_refused():
    with pytest.raises(ce.CryptoError, match="tanınmayan"):
        ce.decrypt_value("ENC1:rot13:QUJD", b"0" * 32)


def test_reveal_does_not_pass_filenames_through_a_shell(tmp_path, monkeypatch):
    """Kabuk metakarakterli dosya adı komut çalıştırmamalı."""
    import app

    calls = []
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: calls.append(a))
    monkeypatch.setattr(os, "startfile", lambda *a: calls.append(a), raising=False)

    evil = str(tmp_path / 'ev"; touch PWNED; echo ".xlsx')
    app.App._reveal(evil)

    assert not (tmp_path / "PWNED").exists()
    for args in calls:
        for arg in args:
            assert isinstance(arg, (list, str))
            assert not isinstance(arg, str) or ";" not in arg or arg == evil
